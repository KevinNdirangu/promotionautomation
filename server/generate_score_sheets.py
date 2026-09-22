import json
import os
import sys
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def shade(cell, color):
    properties = cell._tc.get_or_add_tcPr()
    element = OxmlElement('w:shd')
    element.set(qn('w:fill'), color)
    properties.append(element)

def border(cell):
    properties = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        tag = OxmlElement(f'w:{edge}')
        tag.set(qn('w:val'), 'single')
        tag.set(qn('w:sz'), '6')
        tag.set(qn('w:color'), '808080')
        borders.append(tag)
    properties.append(borders)

def set_cell(cell, text, bold=False, size=12, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ''
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0) # Ensure no space after paragraphs in rubric
    run = paragraph.add_run(str(text or ''))
    run.bold = bold
    run.font.size = Pt(size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    border(cell)

def add_candidate_sheet(doc, candidate, template, settings, crest, first):
    if not first:
        doc.add_page_break()
    if crest and os.path.exists(crest):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.add_run().add_picture(crest, width=Inches(0.6))
        
    for text, size in [('KENYA PRISONS SERVICE', 12), (str(settings.get('station_name', '')).upper() + ' INTERVIEW', 11), (f"{settings.get('promotion_month_name', '')} {settings.get('promotion_year', '')} PROMOTIONAL INTERVIEWS SCORE SHEET", 11)]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(size)
        p.paragraph_format.space_after = Pt(1)
        
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(f"RANK APPLIED FOR: {candidate.get('rank_applied_for', '')}")
    r.bold = True
    r.font.size = Pt(11)
    
    # Profile block as a table
    profile = doc.add_table(rows=3, cols=2)
    profile.alignment = WD_TABLE_ALIGNMENT.CENTER
    fields = [('PF/NO', candidate.get('pf_no')), ('RANK', candidate.get('current_rank')), ('NAME', candidate.get('name')), ('SECTION DEPLOYED', candidate.get('section_deployed')), ('GENDER', candidate.get('gender')), ('ETHNICITY', candidate.get('ethnicity'))]
    for cell, (label, value) in zip([c for row in profile.rows for c in row.cells], fields):
        # We can use set_cell but maybe slightly smaller font so it fits nicely
        set_cell(cell, f'{label}: {value or ""}', bold=True, size=11)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    
    # The rubric table
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ['S/NO', 'VARIABLE', 'MARKS', 'MARKS AWARDED']
    for cell, value in zip(table.rows[0].cells, headers):
        set_cell(cell, value, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(cell, 'D9E2F3')
    for index, criterion in enumerate(template['criteria'], 1):
        cells = table.add_row().cells
        set_cell(cells[0], index, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell(cells[1], str(criterion['label']).upper())
        set_cell(cells[2], criterion['max'], align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell(cells[3], '', align=WD_ALIGN_PARAGRAPH.CENTER)
    cells = table.add_row().cells
    set_cell(cells[0], 'TOTAL MARKS 100%', bold=True)
    set_cell(cells[1], 'TOTAL MARKS 100%', bold=True)
    set_cell(cells[2], '100%', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(cells[3], '', align=WD_ALIGN_PARAGRAPH.CENTER)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # Adding the Appendix at the bottom (below the rubric)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.add_run('APPENDIX 01: EDUCATION').bold = True
    
    edu_items = template.get('educationAppendix', [])
    if edu_items:
        # Borderless table for perfect alignment
        edu_table = doc.add_table(rows=len(edu_items), cols=2)
        edu_table.alignment = WD_TABLE_ALIGNMENT.LEFT
        for i, item in enumerate(edu_items):
            c1, c2 = edu_table.rows[i].cells
            c1.paragraphs[0].paragraph_format.space_after = Pt(0)
            c2.paragraphs[0].paragraph_format.space_after = Pt(0)
            c1.width = Inches(3.0)
            c2.width = Inches(2.0)
            r1 = c1.paragraphs[0].add_run(str(item['qualification']))
            r1.font.size = Pt(11)
            r2 = c2.paragraphs[0].add_run(f"- {item['marks']} MARKS")
            r2.font.size = Pt(11)
            
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(4)
    p.add_run('APPENDIX 02: LENGTH OF SERVICE').bold = True
    
    p2 = doc.add_paragraph(template.get('serviceAppendix', ''))
    p2.paragraph_format.space_after = Pt(0)
    if p2.runs:
        p2.runs[0].font.size = Pt(11)

    p3 = doc.add_paragraph()
    p3.paragraph_format.space_after = Pt(0)
    p3.paragraph_format.space_before = Pt(8)
    p3.add_run('CHAIRMAN SIGNATURE: ____________________\t\tDATE: ____________________').font.size = Pt(10)
    
    for _ in range(4):
        p4 = doc.add_paragraph()
        p4.paragraph_format.space_after = Pt(0)
        p4.paragraph_format.space_before = Pt(2)
        p4.add_run('MEMBER SIGNATURE: ______________________\t\tDATE: ____________________').font.size = Pt(10)

def main():
    with open(sys.argv[1], encoding='utf8') as source: payload = json.load(source)
    settings = payload['settings']
    month_names = ['', 'JANUARY', 'FEBRUARY', 'MARCH', 'APRIL', 'MAY', 'JUNE', 'JULY', 'AUGUST', 'SEPTEMBER', 'OCTOBER', 'NOVEMBER', 'DECEMBER']
    settings['promotion_month_name'] = month_names[int(settings.get('promotion_month') or 1)]
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.40); section.bottom_margin = Inches(0.40); section.left_margin = Inches(0.55); section.right_margin = Inches(0.55)
    for index, candidate in enumerate(payload['candidates']): add_candidate_sheet(doc, candidate, payload['template'], settings, payload['crestPath'], index == 0)
    doc.save(payload['outputPath'])

if __name__ == '__main__': main()
